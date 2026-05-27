"""
doc_i18n.py — Locale helper and runtime patches for Word document builders.

Set DOC_LANG=sq (or call set_lang('sq')) before importing build modules.
"""
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))
_SQ_CACHE = None
_LANG = os.environ.get("DOC_LANG", "en").lower()
_PATCHED = False


def _load_sq():
    global _SQ_CACHE
    if _SQ_CACHE is None:
        path = os.path.join(HERE, "locales", "sq.json")
        with open(path, encoding="utf-8") as f:
            _SQ_CACHE = json.load(f)
    return _SQ_CACHE


def set_lang(lang: str) -> None:
    global _LANG
    _LANG = (lang or "en").lower()


def get_lang() -> str:
    return _LANG


def t(text: str) -> str:
    """Return Albanian when DOC_LANG=sq, else the English source string."""
    if _LANG != "sq" or not text:
        return text
    sq = _load_sq()
    if text in sq:
        return sq[text]
    # add_run often appends "\n" — look up the line without the suffix
    if text.endswith("\n"):
        base = text[:-1]
        if base in sq:
            return sq[base] + "\n"
    return text


def suffix() -> str:
    return "_SQ" if _LANG == "sq" else ""


def _local_filename(name: str) -> str:
    if _LANG != "sq" or name.endswith("_SQ.docx"):
        return name
    return name.replace(".docx", "_SQ.docx")


def _patch_docbuilder():
    from build_all_docs import DocBuilder

    def wrap1(orig):
        def wrapped(self, text, *args, **kwargs):
            if isinstance(text, str):
                text = t(text)
            return orig(self, text, *args, **kwargs)
        return wrapped

    def wrap_bullet(orig):
        def wrapped(self, text, *args, bold_prefix=None, **kwargs):
            if isinstance(bold_prefix, str):
                bold_prefix = t(bold_prefix)
            if isinstance(text, str):
                text = t(text)
            return orig(self, text, *args, bold_prefix=bold_prefix, **kwargs)
        return wrapped

    def wrap_table(orig):
        def wrapped(self, headers, rows, *args, **kwargs):
            headers = [t(h) if isinstance(h, str) else h for h in headers]
            rows = [
                [t(c) if isinstance(c, str) else c for c in row]
                for row in rows
            ]
            return orig(self, headers, rows, *args, **kwargs)
        return wrapped

    def wrap_save(orig):
        def wrapped(self, filename, *args, **kwargs):
            return orig(self, _local_filename(filename), *args, **kwargs)
        return wrapped

    DocBuilder.h = wrap1(DocBuilder.h)
    DocBuilder.body = wrap1(DocBuilder.body)
    DocBuilder.callout = wrap1(DocBuilder.callout)
    DocBuilder.bullet = wrap_bullet(DocBuilder.bullet)
    DocBuilder.table = wrap_table(DocBuilder.table)
    DocBuilder.save = wrap_save(DocBuilder.save)


def _patch_word_doc_module():
    import build_word_doc as bwd

    def wrap_fn(orig):
        def wrapped(text, *args, **kwargs):
            if isinstance(text, str):
                text = t(text)
            return orig(text, *args, **kwargs)
        return wrapped

    def wrap_bullet(orig):
        def wrapped(text, *args, bold_prefix=None, **kwargs):
            if isinstance(bold_prefix, str):
                bold_prefix = t(bold_prefix)
            if isinstance(text, str):
                text = t(text)
            return orig(text, *args, bold_prefix=bold_prefix, **kwargs)
        return wrapped

    def wrap_table(orig):
        def wrapped(headers, rows, *args, **kwargs):
            headers = [t(h) if isinstance(h, str) else h for h in headers]
            rows = [
                [t(c) if isinstance(c, str) else c for c in row]
                for row in rows
            ]
            return orig(headers, rows, *args, **kwargs)
        return wrapped

    bwd.heading = wrap_fn(bwd.heading)
    bwd.body = wrap_fn(bwd.body)
    bwd.bullet = wrap_bullet(bwd.bullet)
    bwd.callout_box = wrap_fn(bwd.callout_box)
    bwd.add_table = wrap_table(bwd.add_table)

    _orig_save = bwd.doc.save

    def localized_save(path, *args, **kwargs):
        base = os.path.basename(path)
        folder = os.path.dirname(path)
        return _orig_save(os.path.join(folder, _local_filename(base)), *args, **kwargs)

    bwd.doc.save = localized_save


_orig_add_run = None


def _patch_cover_runs():
    """Translate all paragraph.add_run text when locale is sq."""
    global _orig_add_run

    def patched_add_run(self, text="", *args, **kwargs):
        if isinstance(text, str) and text and get_lang() == "sq":
            text = t(text)
        return _orig_add_run(self, text, *args, **kwargs)

    from docx.text.paragraph import Paragraph

    if _orig_add_run is None:
        _orig_add_run = Paragraph.add_run
        Paragraph.add_run = patched_add_run


def apply_doc_patches():
    global _PATCHED
    if _PATCHED:
        return
    # add_run must be patched before build_word_doc is imported (module builds on import)
    _patch_cover_runs()
    _patch_docbuilder()
    _patch_word_doc_module()
    _PATCHED = True
